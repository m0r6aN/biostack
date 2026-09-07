from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Pt


SOURCE = Path(r"D:\Repos\BioStack\docs\legal\BioStack_Public_Legal_Policies_v0.9.docx")
OUTPUT = Path(r"D:\Repos\BioStack\docs\legal\BioStack_Public_Legal_Policies_v0.10.docx")


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """Replace visible text while preserving the first run's formatting."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def insert_after(paragraph: Paragraph, text: str, style_name: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if style_name:
        new_paragraph.style = style_name
    elif paragraph.style:
        new_paragraph.style = paragraph.style
    new_paragraph.add_run(text)
    return new_paragraph


def set_cell_text(cell, text: str, *, size_pt: float = 8.5, bold: bool = False) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size_pt)
            run.bold = bold


def remove_table_row(table, row_index: int) -> None:
    row = table.rows[row_index]
    table._tbl.remove(row._tr)


def rewrite_table(table, rows: list[list[str]], *, size_pt: float = 8.5) -> None:
    while len(table.rows) > 1:
        remove_table_row(table, 1)
    for values in rows:
        row = table.add_row()
        for cell, value in zip(row.cells, values, strict=True):
            set_cell_text(cell, value, size_pt=size_pt)


def add_comment(doc, paragraph: Paragraph, text: str) -> None:
    runs = [run for run in paragraph.runs if run.text]
    if runs:
        doc.add_comment(runs=runs, text=text, author="Drafting note", initials="DN")


doc = Document(SOURCE)
original = list(doc.paragraphs)
tables = list(doc.tables)

# Version/date/front-matter updates.
for index in (5, 85, 182, 213, 242):
    replace_paragraph_text(original[index], "Version: 0.10.0-draft")
for index in (6, 86, 183, 214):
    replace_paragraph_text(original[index], "Effective date: August 28, 2026")
replace_paragraph_text(original[243], "Last updated: August 28, 2026")
replace_paragraph_text(
    original[7],
    'Legal operator: Morgan Findings LLC, Georgia ("BioStack," "we," "us," or "our")',
)
confirmed_address = "1800 Twin Branch Dr., Marietta, GA 30062"
replace_paragraph_text(original[8], f"Mailing address: {confirmed_address}")
replace_paragraph_text(original[83], f"Mail: BioStack - {confirmed_address}")
replace_paragraph_text(
    original[87],
    f"Controller and operator: Morgan Findings LLC, Georgia, {confirmed_address}",
)
replace_paragraph_text(original[179], f"Mail: BioStack - {confirmed_address}")
add_comment(
    doc,
    original[8],
    "Operator confirmation required under R7: verify that this is Morgan Findings LLC's current registered mailing address before counsel sign-off and publication.",
)

header = doc.sections[0].header.paragraphs[0]
replace_paragraph_text(header, "BioStack | Public Legal Policies | v0.10 draft")

# Cover metadata table.
set_cell_text(tables[0].cell(0, 1), "0.10.0-draft", size_pt=9.5)
set_cell_text(tables[0].cell(1, 1), "August 28, 2026", size_pt=9.5)
set_cell_text(
    tables[0].cell(2, 1),
    "Operator and licensed-counsel review required; confirmation gates remain open",
    size_pt=9.5,
)
set_cell_text(
    tables[0].cell(3, 1),
    "U.S. adults 18+, hosted account platform, restricted provider pilot",
    size_pt=9.5,
)

# R8: align the Service category without changing the existing non-prescriptive sentences.
replace_paragraph_text(
    original[9],
    'These Terms of Use form a binding agreement between you and BioStack concerning your access to and use of the BioStack websites, free public source-graded evidence library, applications, calculators, knowledge materials, exports, paid tracking and analysis layers, provider-pilot features, and related services (collectively, the "Service"). Read them carefully.',
)
terms_service_paragraph = original[17].insert_paragraph_before(
    "BioStack provides a free, public, source-graded evidence library, with optional paid tracking and analysis layers.",
    style="Normal",
)
terms_service_paragraph.paragraph_format.space_after = Pt(6)
privacy_service_paragraph = original[88].insert_paragraph_before(
    "BioStack provides a free, public, source-graded evidence library, with optional paid tracking and analysis layers.",
    style="Normal",
)
privacy_service_paragraph.paragraph_format.space_after = Pt(6)

# R1: replace the local-first premise with hosted custody throughout the Terms.
replace_paragraph_text(
    original[35],
    "BioStack stores account and user-entered Service data on hosted systems. You remain responsible for device security and for securely storing or deleting any exports or copies you create. Loss of a device may expose local copies, but the authoritative Service records associated with your account are maintained by BioStack and its approved processors according to the Privacy Policy.",
)
replace_paragraph_text(
    original[55],
    "Our Privacy and Consumer Health Data Policy explains what information we process, why, how long we retain it, when we disclose it, and how to exercise rights. BioStack can service verified access, export, correction, and deletion requests for account data it maintains, subject to the lawful retention exceptions described in that Policy.",
)
replace_paragraph_text(
    original[60],
    "The Service may include beta, preview, or pilot features. They may change, fail, produce incomplete results, or be withdrawn. Do not rely on availability, compatibility, data preservation, or continued support unless a signed agreement states otherwise.",
)

# R9: provider-pilot intake consent is versioned and durable.
pilot_note = insert_after(
    original[53],
    "A provider access request records the consent version presented and the time consent is recorded. That versioned consent record, together with the normalized email and request metadata, is the acceptance record for provider-pilot intake. Submitting a request does not create access or authorize use; requests remain subject to review, honeypot and rate-limiting controls, and any required signed pilot agreement.",
    "Normal",
)

# R3: insert monthly-only billing terms and renumber the following Terms sections.
billing_heading = original[62].insert_paragraph_before(
    "15. Billing, renewal, cancellation, and refunds", style="Heading 2"
)
billing_paragraphs = [
    "Observer is free. Operator costs USD $12 per month, and Commander costs USD $29 per month. BioStack currently offers monthly billing only; annual billing is not offered. Paid subscriptions renew automatically each month until canceled.",
    "Stripe processes paid subscriptions and payment information under its own terms and privacy policy. BioStack does not receive or store full payment-card numbers. You authorize BioStack and Stripe to charge the disclosed monthly amount and applicable taxes using your selected payment method.",
    "You may cancel through the available billing controls. Unless applicable law requires otherwise or BioStack expressly states otherwise, cancellation takes effect at the end of the then-current monthly billing period, fees already charged are non-refundable, and BioStack does not provide prorated refunds or credits for partial months.",
    "If payment fails or a paid subscription is no longer in an active or trialing state, paid access fails closed and the account immediately reverts to the Observer tier without a grace period. BioStack may retry payment through Stripe and will restore paid access only after the subscription returns to an eligible state. Your content remains subject to the Privacy Policy and its retention and deletion terms.",
]
for text in billing_paragraphs:
    original[62].insert_paragraph_before(text, style="Normal")
add_comment(
    doc,
    billing_heading,
    "Licensed counsel must confirm the cancellation and refund language for each launch jurisdiction before publication.",
)
for index, number in zip((62, 65, 68, 72, 74, 77, 79), range(16, 23), strict=True):
    heading_text = original[index].text
    suffix = heading_text.split(". ", 1)[1]
    replace_paragraph_text(original[index], f"{number}. {suffix}")

# R1: Privacy Policy hosted-custody rewrite.
replace_paragraph_text(original[89], "1. Scope and hosted-data custody")
replace_paragraph_text(
    original[90],
    "This Policy applies to BioStack websites, applications, support channels, pilots, account features, and other hosted features that link to it.",
)
replace_paragraph_text(
    original[91],
    "BioStack's running Service uses server-side persistence. When you create an account or enter, save, import, sync, submit, or otherwise provide information through an enabled Service feature, BioStack and approved processors receive and store that information on hosted systems as described in this Policy.",
)
replace_paragraph_text(
    original[92],
    "Information contained only in a copy or export you create outside the Service is controlled by you. The corresponding account data stored by the Service remains within BioStack's custody until it is deleted or de-identified under this Policy, subject to lawful retention exceptions.",
)
original[91].paragraph_format.space_after = Pt(6)
original[92].paragraph_format.space_after = Pt(6)
replace_paragraph_text(
    original[100],
    "Values entered into a calculator, selected units, formula version, and resulting mathematical output. Depending on the calculator and whether you save a result, these values may be processed in the browser, transmitted to a hosted API, or stored with your account records. The interface and collection notice should identify any saved or transmitted data.",
)
replace_paragraph_text(
    original[120],
    "BioStack may collect the consumer health data categories described in Section 2 when you deliberately enter, import, save, sync, or submit them through an enabled feature.",
)

# Update the purposes table to remove the local-processing premise.
set_cell_text(
    tables[1].cell(2, 2),
    "User request; hosted processing or storage when transmitted or saved",
    size_pt=8.5,
)

# R2/R6: retention schedule, audit-spine carve-out, and backup restoration behavior.
retention_rows = [
    ["Primary account profile", "While active; delete within 30 days after a valid deletion request or account closure, subject to lawful retention exceptions"],
    ["Primary consumer health and protocol data", "While the account or feature remains enabled; delete from primary systems within 30 days after a valid request"],
    ["Production PostgreSQL PITR backups", "Age out within the configured 7-35 day point-in-time-recovery window; exact setting must be operator-confirmed; data deleted from primary systems is re-deleted if a backup is restored"],
    ["Weekly logical database backups", "90-day lifecycle under the current operations runbook; data deleted from primary systems is re-deleted if a backup is restored"],
    ["Decision Receipts and Governed Spine audit records", "Permanent to preserve the append-only hash chain; access-restricted and pseudonymized where feasible; no product analytics, advertising, or unrelated use"],
    ["Security and access logs", "12 months unless a documented investigation or legal duty requires longer"],
    ["Support records", "24 months after ticket closure, with sensitive attachments removed sooner when no longer needed"],
    ["Privacy and deletion request records", "24 months after closure, excluding deleted health content"],
    ["Policy acceptance, age attestation, and consent evidence", "6 years after the relationship ends, limited to evidence of the event and not the underlying health content"],
    ["Provider-pilot administration and audit records", "6 years after pilot closure, excluding test health content unless a signed schedule requires otherwise"],
    ["Stripe billing and transaction records", "Under Stripe's retention terms and BioStack's applicable financial, tax, dispute, and legal obligations; counsel and operator to confirm"],
    ["Financial and tax records", "As required by applicable law, ordinarily 7 years"],
]
rewrite_table(tables[2], retention_rows, size_pt=8.0)
replace_paragraph_text(
    original[137],
    "The operator must confirm this schedule against production PostgreSQL settings, backup lifecycle policies, Stripe terms, and applicable limitation periods before approval. BioStack may retain a minimal suppression record, transaction record, legal-hold copy, or integrity/audit record where reasonably necessary and permitted by law to comply with law, prevent fraud, preserve audit-chain integrity, establish legal rights, or honor a request not to contact you. Such records are access-restricted and are not used for product analytics, advertising, or unrelated purposes.",
)
retention_carveout = insert_after(
    original[137],
    "On a verified deletion request, BioStack deletes personal content from primary systems and instructs applicable processors to do the same. Integrity and audit records, including append-only, hash-chained Decision Receipts, may be retained in restricted and pseudonymized-where-feasible form when reasonably necessary and permitted by law to preserve chain integrity, comply with a legal hold or legal duty, prevent fraud, or establish, exercise, or defend legal rights. Licensed counsel must confirm this treatment under applicable consumer health data laws before publication.",
    "Normal",
)
replace_paragraph_text(
    original[152],
    "BioStack is responsible for servicing verified access, export, correction, and deletion requests for data held in its primary systems and will instruct applicable processors as described in this Policy. You remain responsible for copies or exports you keep outside the Service. Deletion rights are subject only to lawful and proportionate retention exceptions described in Sections 8 and 10.",
)
replace_paragraph_text(
    original[155],
    "Deletion removes or irreversibly de-identifies in-scope personal content from primary systems and instructs applicable processors to do the same. Restricted integrity and audit records, including append-only, hash-chained Decision Receipts, may be retained and pseudonymized where feasible under Section 8 when reasonably necessary and permitted by law. Deletion does not require BioStack to remove public research sources, information that does not identify you, records held solely by an independent recipient, or other minimal records lawfully retained under Section 8.",
)

# R5: real application cookie identified in code; production capture remains a gate.
replace_paragraph_text(
    original[226],
    "The current application code identifies the register below. Before publication, the operator must capture a production session and confirm the exact cookie name, duration, Secure/HttpOnly/SameSite attributes, and absence of additional cookies or local-storage keys.",
)
cookie_rows = [[
    "biostack_session",
    "BioStack",
    "Passwordless authentication and session integrity",
    "Strictly necessary; Secure, HttpOnly, SameSite=Lax in production",
    "30 days; persistent; no sliding expiration",
]]
rewrite_table(tables[3], cookie_rows, size_pt=8.0)
add_comment(
    doc,
    tables[3].cell(1, 0).paragraphs[0],
    "R5 production confirmation gate: capture the live Set-Cookie header and verify this register before counsel sign-off and publication.",
)
replace_paragraph_text(
    original[223],
    "preserve user-requested session state or other strictly necessary device-side state.",
)

# R4: replace the subprocessor placeholder with the current provider register.
replace_paragraph_text(
    original[245],
    "BioStack uses the providers listed below to operate current hosted features. No advertising or analytics subprocessor is authorized.",
)
replace_paragraph_text(
    original[246],
    "GitHub is used for source hosting and continuous integration and is not authorized to receive BioStack user data.",
)
subprocessor_rows = [
    [
        "Microsoft Corporation (Microsoft Azure)",
        "Application hosting and PostgreSQL database",
        "Account, contact, consumer health/protocol, calculator, consent, technical, security, and audit data",
        "United States; configured Azure regions to be operator-confirmed",
        "Not applicable to the U.S.-only launch; counsel to confirm",
        "Per Sections 8-10; primary deletion instructions and backup aging apply",
        "August 28, 2026 (draft)",
    ],
    [
        "Stripe, Inc. (contracting entity to confirm)",
        "Payment and subscription processing",
        "Billing contact, customer/subscription identifiers, transaction status, and payment data handled by Stripe",
        "United States and other locations under Stripe's terms",
        "Not applicable to the U.S.-only launch; counsel to confirm",
        "Under Stripe's terms and applicable financial/tax law; operator and counsel to confirm",
        "August 28, 2026 (draft)",
    ],
    [
        "[OPERATOR TO CONFIRM LEGAL NAME OF EMAIL RELAY]",
        "Magic-link and transactional email delivery",
        "Email address, magic-link message, delivery metadata, and transactional content",
        "Operator to confirm",
        "Operator and counsel to confirm",
        "Transient delivery and log retention; operator to confirm",
        "August 28, 2026 (draft)",
    ],
    [
        "GitHub, Inc.",
        "Source hosting and continuous integration only",
        "No BioStack user data authorized",
        "United States/global infrastructure",
        "Not applicable; no user data authorized",
        "Repository and CI retention only; no user-data retention",
        "August 28, 2026 (draft)",
    ],
]
rewrite_table(tables[4], subprocessor_rows, size_pt=7.5)
add_comment(
    doc,
    tables[4].cell(3, 0).paragraphs[0],
    "R4 operator confirmation gate: identify the current production email provider and its legal contracting entity before counsel sign-off and publication.",
)

# Core properties and final document write.
doc.core_properties.title = "BioStack Public Legal Policies v0.10"
doc.core_properties.subject = "Terms, privacy, consumer health data, calculator safety, cookies, and subprocessors"
doc.core_properties.comments = "Working draft generated from POLICY-REDLINE-v0.9-to-v0.10.md; operator and licensed-counsel review required."
doc.core_properties.modified = datetime.now(timezone.utc)
doc.save(OUTPUT)
print(OUTPUT)
